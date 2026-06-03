import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── colour palette ──────────────────────────────────────────────────────────
HEADER_BG   = "1F4E79"   # dark blue
HEADER_FG   = "FFFFFF"   # white
SECTION_BG  = "2E75B6"   # medium blue
SECTION_FG  = "FFFFFF"
REQ_BG      = "FFF2CC"   # light yellow  – required
OPT_BG      = "EBF3E8"   # light green   – optional
ALT_REQ_BG  = "FFEB9C"   # slightly darker yellow for alternating required rows
ALT_OPT_BG  = "DDEEDD"   # slightly darker green  for alternating optional rows

def hex_fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def make_header(ws, row, col, value, bold=True):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font      = Font(name="Calibri", bold=bold, color=HEADER_FG, size=11)
    cell.fill      = hex_fill(HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = thin_border()
    return cell

def make_section_header(ws, row, col_start, col_end, title):
    ws.merge_cells(start_row=row, start_column=col_start,
                   end_row=row,   end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=title)
    cell.font      = Font(name="Calibri", bold=True, color=SECTION_FG, size=12)
    cell.fill      = hex_fill(SECTION_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin_border()
    return cell

def make_data_cell(ws, row, col, value, required, row_index):
    cell = ws.cell(row=row, column=col, value=value)
    if required:
        bg = REQ_BG if row_index % 2 == 0 else ALT_REQ_BG
    else:
        bg = OPT_BG if row_index % 2 == 0 else ALT_OPT_BG
    cell.fill      = hex_fill(bg)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border    = thin_border()
    cell.font      = Font(name="Calibri", size=10)
    return cell

def make_badge_cell(ws, row, col, required):
    text  = "Obligatoire" if required else "Optionnel"
    color = "C00000" if required else "375623"    # dark red / dark green
    cell  = ws.cell(row=row, column=col, value=text)
    cell.font      = Font(name="Calibri", bold=True, color=color, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin_border()
    # keep same row background
    req_idx = 0  # just pick base colour for badge column
    if required:
        cell.fill = hex_fill(REQ_BG)
    else:
        cell.fill = hex_fill(OPT_BG)
    return cell

# ── data definition ─────────────────────────────────────────────────────────
# Each entry: (Champ UI, Nom technique, Type, Obligatoire, Description)
SECTIONS = [
    {
        "title": "Informations Personnelles (Personal Info)",
        "fields": [
            ("Matricule",         "badge_id",             "Texte (50 car.)",   False, "Identifiant unique de l'employé, généré automatiquement"),
            ("Prénom",            "employee_first_name",  "Texte (200 car.)",  True,  "Prénom de l'employé"),
            ("Nom",               "employee_last_name",   "Texte (200 car.)",  False, "Nom de famille de l'employé"),
            ("Email",             "email",                "Email (254 car.)",  True,  "Adresse e-mail unique de l'employé"),
            ("Téléphone",         "phone",                "Texte (15 car.)",   True,  "Numéro de téléphone de l'employé"),
            ("Date de naissance", "dob",                  "Date",              False, "Date de naissance (format JJ/MM/AAAA)"),
            ("Genre",             "gender",               "Choix",             False, "Homme / Femme (défaut : Homme)"),
            ("CIN",               "qualification",        "Texte (50 car.)",   False, "Numéro de la Carte Nationale d'Identité"),
            ("Photo de profil",   "employee_profile",     "Image",             False, "Photo de profil de l'employé"),
        ],
    },
    {
        "title": "Informations Professionnelles (Work Info)",
        "fields": [
            ("Poste",                   "job_position_id",       "FK → JobPosition",   False, "Poste occupé par l'employé"),
            ("Département",             "department_id",         "FK → Department",    False, "Département d'appartenance"),
            ("Type de travail",         "work_type_id",          "FK → WorkType",      False, "Type de contrat / travail"),
            ("Type d'employé",          "employee_type_id",      "FK → EmployeeType",  False, "Catégorie d'employé"),
            ("Rôle",                    "job_role_id",           "FK → JobRole",       False, "Rôle dans l'organisation"),
            ("Responsable hiérarchique","reporting_manager_id",  "FK → Employee",      False, "Manager direct"),
            ("Société",                 "company_id",            "FK → Company",       False, "Société employeuse"),
            ("Lieu de travail",         "location",              "Texte (50 car.)",    False, "Localisation du poste"),
            ("Équipe / Shift",          "shift_id",              "FK → EmployeeShift", False, "Horaire / équipe de travail"),
            ("Date d'embauche",         "date_joining",          "Date",               False, "Date de début de contrat"),
            ("Date fin de contrat",     "contract_end_date",     "Date",               False, "Date de fin de contrat (si applicable)"),
            ("Salaire de base",         "basic_salary",          "Entier (DH)",        False, "Salaire de base mensuel en dirhams"),
            ("Salaire horaire",         "salary_hour",           "Entier (DH)",        False, "Taux horaire"),
            ("Numéro CNSS",             "email (work)",          "Texte (254 car.)",   False, "Numéro d'affiliation CNSS"),
            ("Affiliation CIMR",        "mobile",                "Texte (254 car.)",   False, "Numéro d'affiliation CIMR"),
            ("Numéro CIMR",             "numero_cimr",           "Texte (254 car.)",   False, "Numéro CIMR de l'employé"),
            ("Affiliation MUT",         "affiliation_mut",       "Texte (254 car.)",   False, "Affiliation mutuelle"),
            ("Tags",                    "tags",                  "M2M → EmployeeTag",  False, "Étiquettes / catégories libres"),
        ],
    },
    {
        "title": "Informations Bancaires (Bank Details)",
        "fields": [
            ("Nom de la banque",   "bank_name",       "Texte (50 car.)",  False, "Nom de l'établissement bancaire"),
            ("Numéro de compte",   "account_number",  "Texte (50 car.)",  False, "Numéro de compte bancaire"),
            ("Agence / Branche",   "branch",          "Texte (50 car.)",  False, "Nom ou code de l'agence"),
            ("Adresse banque",     "address",         "Texte (255 car.)", False, "Adresse de l'agence bancaire"),
            ("Pays",               "country",         "Texte (50 car.)",  False, "Pays de la banque"),
            ("Région / État",      "state",           "Texte (50 car.)",  False, "Région ou état"),
            ("Ville",              "city",            "Texte (50 car.)",  False, "Ville de l'agence"),
            ("Code 1",             "any_other_code1", "Texte (50 car.)",  False, "Code complémentaire 1 (ex. code SWIFT)"),
            ("Code 2",             "any_other_code2", "Texte (50 car.)",  False, "Code complémentaire 2 (ex. IBAN)"),
        ],
    },
]

# ── build worksheet ─────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Champs Employé"

# column widths
COL_WIDTHS = [4, 28, 28, 22, 16, 52]
for i, w in enumerate(COL_WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 40
ws.row_dimensions[2].height = 22

# Title row
ws.merge_cells("A1:F1")
title_cell = ws["A1"]
title_cell.value     = "Fiche des Champs – Ajout Employé (FlowRH)"
title_cell.font      = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
title_cell.fill      = hex_fill("0D2137")
title_cell.alignment = Alignment(horizontal="center", vertical="center")

# Column headers
HEADERS = ["#", "Champ (Interface)", "Nom Technique", "Type", "Obligatoire", "Description"]
for col, h in enumerate(HEADERS, 1):
    make_header(ws, 2, col, h)

current_row = 3
field_number = 1

for section in SECTIONS:
    # section title row
    make_section_header(ws, current_row, 1, 6, section["title"])
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    for idx, (ui_label, tech_name, ftype, required, desc) in enumerate(section["fields"]):
        ws.row_dimensions[current_row].height = 18

        # #
        num_cell = ws.cell(row=current_row, column=1, value=field_number)
        num_cell.font      = Font(name="Calibri", bold=True, size=10,
                                  color="FFFFFF" if required else "375623")
        num_cell.fill      = hex_fill("C00000" if required else "375623")
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border    = thin_border()

        make_data_cell(ws, current_row, 2, ui_label,  required, idx)
        make_data_cell(ws, current_row, 3, tech_name, required, idx)
        make_data_cell(ws, current_row, 4, ftype,     required, idx)
        make_badge_cell(ws, current_row, 5, required)
        make_data_cell(ws, current_row, 6, desc,      required, idx)

        current_row  += 1
        field_number += 1

# ── legend sheet ─────────────────────────────────────────────────────────────
lg = wb.create_sheet("Légende")
lg.column_dimensions["A"].width = 22
lg.column_dimensions["B"].width = 52

lg.merge_cells("A1:B1")
t = lg["A1"]
t.value     = "Légende"
t.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
t.fill      = hex_fill("0D2137")
t.alignment = Alignment(horizontal="center", vertical="center")
lg.row_dimensions[1].height = 28

items = [
    (REQ_BG, "C00000", "Champ OBLIGATOIRE – doit être renseigné pour créer l'employé"),
    (OPT_BG, "375623", "Champ OPTIONNEL – peut être laissé vide"),
    ("FFFFFF", "1F4E79", "FK = clé étrangère (référence à une autre table)"),
    ("FFFFFF", "1F4E79", "M2M = relation plusieurs-à-plusieurs"),
]
for r, (bg, fg, text) in enumerate(items, 2):
    swatch = lg.cell(row=r, column=1, value="■ " + ("Obligatoire" if "OBLIG" in text else "Optionnel" if "OPTION" in text else "Info"))
    swatch.font      = Font(name="Calibri", bold=True, color=fg, size=11)
    swatch.fill      = hex_fill(bg)
    swatch.alignment = Alignment(horizontal="center", vertical="center")
    swatch.border    = thin_border()
    desc_cell        = lg.cell(row=r, column=2, value=text)
    desc_cell.font   = Font(name="Calibri", size=10)
    desc_cell.border = thin_border()
    desc_cell.alignment = Alignment(vertical="center", wrap_text=True)
    lg.row_dimensions[r].height = 22

# ── freeze panes & auto-filter ────────────────────────────────────────────
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:F2"

# ── save ─────────────────────────────────────────────────────────────────────
OUTPUT = "/home/abdellah/FlowRH/employee_fields.xlsx"
wb.save(OUTPUT)
print(f"Fichier créé : {OUTPUT}")
